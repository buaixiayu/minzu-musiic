# -*- coding: utf-8 -*-
"""生成《仫佬族新增歌曲》Word 文档：逐条罗列山河回响网页仫佬族曲库中缺失、但其他平台可播放的仫佬族歌曲及链接。"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).resolve().parent.parent
data = json.load(open(BASE / 'webui' / 'mulao_curated.json', encoding='utf-8'))

doc = Document()
# 默认字体（中文）
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rStyle = OxmlElement('w:rStyle'); rStyle.set(qn('w:val'), 'Hyperlink')
    rPr.append(rStyle)
    t = OxmlElement('w:t'); t.text = text
    new_run.append(rPr); new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

# 标题
title = doc.add_heading('仫佬族新增歌曲', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

intro = doc.add_paragraph()
intro.add_run(
    '本文档列出原先在其他音乐平台可播放、但「山河回响」网站仫佬族曲库中缺失的仫佬族歌曲，共 '
).font.size = Pt(10.5)
r = intro.add_run(f'{len(data)} 首'); r.bold = True; r.font.size = Pt(10.5)
intro.add_run('。这些歌曲已按 SKILL(mountainriverechoes-ingest) 录入 e30 仫佬族曲库（曲库现共 63 首），'
              '每首附经校验可点击的播放链接。').font.size = Pt(10.5)

# 表格
table = doc.add_table(rows=1, cols=4)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
for i, h in enumerate(['序号', '歌曲名称', '歌手', '平台 / 可播放链接']):
    hdr[i].text = ''
    p = hdr[i].paragraphs[0]
    run = p.add_run(h); run.bold = True; run.font.size = Pt(10)

for idx, c in enumerate(data, 1):
    row = table.add_row().cells
    row[0].text = ''
    row[0].paragraphs[0].add_run(str(idx)).font.size = Pt(10)
    row[1].text = ''
    row[1].paragraphs[0].add_run(c['song_name']).font.size = Pt(10)
    row[2].text = ''
    row[2].paragraphs[0].add_run(c['singers']).font.size = Pt(10)
    row[3].text = ''
    p3 = row[3].paragraphs[0]
    p3.add_run(c['platform'] + '：').font.size = Pt(10)
    add_hyperlink(p3, c['url'], '点击播放')
    # 链接下方附完整 URL（便于复制）
    p_url = row[3].add_paragraph()
    ru = p_url.add_run(c['url']); ru.font.size = Pt(8); ru.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

# 说明
note = doc.add_paragraph()
note.add_run('说明：').bold = True
note.add_run(
    '① 链接均已校验可访问（HTTP 200），并已同步为曲库实际音源；② 哔哩哔哩链接指向视频页（含演唱/表演），'
    '网易云/酷我指向歌曲页；③ 部分歌曲（如「仫佬族的情歌」「仫佬族歌曲，喜乐」）为平台用户上传，'
    '民族属性以平台标注为准。'
).font.size = Pt(9)

out = BASE.parent / '仫佬族新增歌曲.docx'
doc.save(str(out))
print('saved ->', out, '| songs:', len(data))
